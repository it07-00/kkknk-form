import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('Bảng cung cấp số liệu khí nhà kính.xlsx', data_only=True)
print('Sheet names:', wb.sheetnames)

for name in wb.sheetnames:
    sheet = wb[name]
    print(f'\n======================================================')
    print(f'SHEET: {name} (Rows: {sheet.max_row}, Cols: {sheet.max_column})')
    print(f'======================================================')
    for r in range(1, sheet.max_row + 1):
        row_vals = [sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)]
        # Filter trailing Nones
        while row_vals and row_vals[-1] is None:
            row_vals.pop()
        if row_vals:
            # print non-empty rows
            formatted = [str(v).strip() if v is not None else '' for v in row_vals]
            print(f'Row {r:3d}: {" | ".join(formatted)}')
